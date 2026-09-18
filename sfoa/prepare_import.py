"""Prepare a deduplicated, auditable Salesforce Bulk API CSV from Jianyu exports.

This script never connects to Salesforce. The output contains commercial data and
stays in the ignored local upload directory, never in Git.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import warnings
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

FIELDS = [
    "Source_Key__c", "Source_Row_Number__c", "Name", "Full_Title__c", "Project_Name__c",
    "Matching_Keywords__c", "bidstate__c", "bidcity__c", "Source_District__c",
    "Publication_Category__c", "type__c", "Tender_Stage__c", "release_date__c",
    "Source_Industry__c", "Buyer_Name__c", "Winner_Name__c",
    "Buyer_Type__c", "Buyer_Contact__c", "Buyer_Phone__c",
    "Buyer_Address__c", "Agency_Name__c",
    "Winner_Notice_Contact__c", "Winner_Notice_Phone__c",
    "Winner_Registry_Contact__c", "Winner_Registry_Phone__c",
    "Winner_Registry_Email__c",
    "accountname__c", "bidwinningname__c",
    "Full_Project_Number__c", "projectnumber__c",
    "Source_Budget_Wan__c", "Source_Winning_Wan__c",
    "Source_Content__c", "content__c", "Source_Content_Hash__c",
    "Project_Scope__c", "Registration_Deadline__c", "bidopeningdate__c",
    "biddeadline__c", "Contract_Signed_Date__c",
    "website__c", "Announcement_Url_Full__c", "swordfishwebsite__c", "Source_File__c",
]
AWARDED = {"中标", "成交"}
OPEN = {"招标", "邀标", "竞价", "竞谈", "询价", "磋商", "单一", "邀请",
        "变更", "采购意向", "需求公示", "预告"}
ALLOWED_CATEGORIES = frozenset({
    "招标", "邀标", "询价", "竞谈", "单一", "竞价", "变更", "中标", "成交",
})
HYPERLINK = re.compile(
    # Jianyu sometimes puts unescaped quotation marks in the visible title.
    # The target URL has no quotes; take the rest through the final quote as
    # display text rather than trying to evaluate the malformed formula.
    r'^HYPERLINK\(\s*"([^"]*)"\s*[,;]\s*"(.*)"\s*\)$',
    re.IGNORECASE | re.DOTALL,
)


def excel_value(cell, location: str):
    """Read the visible text of Jianyu's HYPERLINK formulas without evaluating Excel code."""
    value = cell.value
    if cell.data_type != "f":
        return value
    formula = str(value).lstrip("=").strip()
    match = HYPERLINK.fullmatch(formula)
    if not match:
        raise ValueError(f"无法解析公式，停止上传：{location}")
    target, display = (part.replace('""', '"') for part in match.groups())
    return display or target


def clean(value) -> str:
    return str(value).strip() if value is not None else ""


def day(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = clean(value)
    match = re.search(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", text)
    if not match:
        return ""
    try:
        return datetime.strptime(match.group().replace("/", "-"), "%Y-%m-%d").date().isoformat()
    except ValueError:
        return ""


def amount(value, errors: Counter) -> str:
    text = clean(value).replace(",", "").replace("￥", "").replace("¥", "")
    if not text:
        return ""
    try:
        number = Decimal(text)
        if not number.is_finite() or abs(number) >= Decimal("1000000000000000"):
            raise InvalidOperation
        return str(number)
    except (InvalidOperation, ValueError):
        errors["unparsed_amount"] += 1
        return ""


def source_key(row) -> str:
    # Keep the existing key independent of hyperlink columns so corrected
    # formula parsing updates the same Salesforce records without duplicates.
    parts = [clean(row[i]) for i in (2, 3, 6, 7, 11, 13)]
    parts.insert(0, day(row[8]))
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()


def record(row, file_name: str, errors: Counter) -> dict:
    content = clean(row[7])
    title = clean(row[5]) or clean(row[11]) or content[:80] or "未命名标讯"
    category = clean(row[6])
    stage = "已成交" if category in AWARDED else "招标中" if category in OPEN else "其他"
    result = {
        "Source_Key__c": source_key(row),
        "Source_Row_Number__c": clean(row[0])[:20],
        "Name": title[:80],
        "Full_Title__c": title[:255],
        "Project_Name__c": clean(row[11])[:255],
        "Matching_Keywords__c": clean(row[1])[:255],
        "bidstate__c": clean(row[2])[:10],
        "bidcity__c": clean(row[3])[:10],
        "Source_District__c": clean(row[4])[:80],
        "Publication_Category__c": category[:80],
        "type__c": category if len(category) <= 6 else "",
        "Tender_Stage__c": stage,
        "release_date__c": day(row[8]),
        "Source_Industry__c": clean(row[12])[:255],
        "Buyer_Name__c": clean(row[21])[:255],
        "Winner_Name__c": clean(row[27])[:255],
        "Buyer_Type__c": clean(row[22])[:100],
        "Buyer_Contact__c": clean(row[23])[:100],
        "Buyer_Phone__c": clean(row[24])[:100],
        "Buyer_Address__c": clean(row[25])[:255],
        "Agency_Name__c": clean(row[26])[:255],
        "Winner_Notice_Contact__c": clean(row[28])[:100],
        "Winner_Notice_Phone__c": clean(row[29])[:100],
        "Winner_Registry_Contact__c": clean(row[30])[:100],
        "Winner_Registry_Phone__c": clean(row[31])[:255],
        "Winner_Registry_Email__c": clean(row[32]),
        "accountname__c": clean(row[21]) if len(clean(row[21])) <= 30 else "",
        "bidwinningname__c": clean(row[27]) if len(clean(row[27])) <= 30 else "",
        "Full_Project_Number__c": clean(row[13])[:100],
        "projectnumber__c": clean(row[13]) if len(clean(row[13])) <= 20 else "",
        "Source_Budget_Wan__c": amount(row[15], errors),
        "Source_Winning_Wan__c": amount(row[16], errors),
        "Source_Content__c": content,
        "content__c": content[:32768],
        "Source_Content_Hash__c": hashlib.sha256(content.encode("utf-8")).hexdigest(),
        "Project_Scope__c": clean(row[14]),
        "Registration_Deadline__c": day(row[17]),
        "bidopeningdate__c": day(row[18]),
        "biddeadline__c": day(row[19]),
        "Contract_Signed_Date__c": day(row[20]),
        "website__c": clean(row[9]) if len(clean(row[9])) <= 255 else "",
        "Announcement_Url_Full__c": clean(row[9]),
        "swordfishwebsite__c": clean(row[10]),
        "Source_File__c": file_name[:255],
    }
    if len(content) > 131072:
        errors["content_too_long"] += 1
        raise ValueError("公告正文超过 Salesforce 字段上限；停止而非截断")
    if not result["release_date__c"] or not content:
        errors["missing_key_data"] += 1
        raise ValueError("公告缺少发布日期或正文；停止而非上传不完整记录")
    if len(clean(row[13])) > 100:
        errors["project_number_truncated"] += 1
    return result


def prepare(root: Path, output: Path, selected_files=None) -> dict:
    files = sorted(selected_files if selected_files is not None else
                   (p for p in root.rglob("*.xlsx") if not p.name.startswith("~$")))
    if not files:
        raise FileNotFoundError(f"没有找到 XLSX：{root}")
    output.mkdir(parents=True, exist_ok=True)
    csv_path = output / "bidnews_upsert.csv"
    errors: Counter = Counter()
    by_date: Counter = Counter()
    seen: set[str] = set()
    rows_read = rows_written = duplicates = 0
    categories_skipped: Counter = Counter()
    with csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS, lineterminator="\r\n")
        writer.writeheader()
        for path in files:
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Workbook contains no default style")
                book = load_workbook(path, read_only=True, data_only=False)
            try:
                sheet = book.active
                rows = sheet.iter_rows()
                header = tuple(cell.value for cell in next(rows, ()))
                if len(header) != 33 or clean(header[8]) != "发布时间":
                    raise ValueError(f"文件表头不符合已核验的 33 列格式：{path}")
                next(rows, None)  # second header row
                for row_number, cells in enumerate(rows, start=3):
                    if not cells or not any(cell.value is not None for cell in cells):
                        continue
                    rows_read += 1
                    if len(cells) != 33:
                        raise ValueError(f"数据行列数异常：{path} 第 {row_number} 行")
                    row = tuple(excel_value(cell, f"{path.name} 第{row_number}行 第{index}列")
                                for index, cell in enumerate(cells, start=1))
                    category = clean(row[6])
                    if category not in ALLOWED_CATEGORIES:
                        categories_skipped[category or "(空白)"] += 1
                        continue
                    item = record(row, path.name, errors)
                    if item["Source_Key__c"] in seen:
                        duplicates += 1
                        continue
                    seen.add(item["Source_Key__c"])
                    writer.writerow(item)
                    rows_written += 1
                    by_date[item["release_date__c"]] += 1
            finally:
                book.close()
    report = {
        "source_files": len(files), "rows_read": rows_read,
        "rows_written": rows_written, "duplicates_skipped": duplicates,
        "categories_skipped": dict(categories_skipped),
        "allowed_categories": sorted(ALLOWED_CATEGORIES),
        "warnings": dict(errors), "by_date": dict(by_date),
        "csv": str(csv_path), "csv_sha256": hashlib.sha256(csv_path.read_bytes()).hexdigest(),
    }
    (output / "prepare_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(json.dumps(prepare(args.source, args.output), ensure_ascii=False, indent=2))
